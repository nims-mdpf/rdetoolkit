"""Mode-specific artifact stage order (Session I-REVIEW-A, ruling #5 / review R3).

``InvoiceService.apply_config`` ran structured -> magic -> description for every
mode, and the mode seam was a ``frozenset`` — a set cannot express an order.
contracts.md claimed the difference was unobservable because the three steps
touch disjoint state. Review R3 falsified that: when the framework's own
magic-variable expansion fails, the *order* decides which artifacts a failed
tile leaves behind.

v1's pipelines (``processing/factories.py``):

* invoice / SmartTable: Thumbnail -> StructuredInvoiceSaver -> VariableApplier
  -> DescriptionUpdater
* MultiDataTile / ExcelInvoice: VariableApplier -> Thumbnail ->
  StructuredInvoiceSaver -> DescriptionUpdater
* RDEFormat: Thumbnail -> DescriptionUpdater (no structured export, no magic)

EP table:
| TC | Class | Input | Expected |
|----|-------|-------|----------|
| TC-IRA-5-EP-030 | handler order | 5 modes | the v1 stage sequence, ordered |
| TC-IRA-5-EP-031 | magic failure | multidatatile | v1 oracle: no ``structured/invoice.json`` |
| TC-IRA-5-EP-032 | magic failure | excelinvoice | v1 oracle: no ``structured/invoice.json`` |
| TC-IRA-5-EP-033 | magic failure | invoice | v1 oracle: ``structured/invoice.json`` present |
| TC-IRA-5-EP-036 | structured failure | multidatatile | v1 oracle: magic already expanded in the tile invoice |
| TC-IRA-5-EP-037 | structured failure | excelinvoice | v1 oracle: magic already expanded in the tile invoice |
| TC-IRA-5-EP-038 | structured failure | invoice | v1 oracle: magic never ran (structured comes first) |

BV / negative table:
| TC | Class | Input | Expected |
|----|-------|-------|----------|
| TC-IRA-5-EV-034 | API | ``modes.protocol`` | ``invoice_stage_steps`` is gone |
| TC-IRA-5-EV-035 | default | no provider | the invoice order is the default |
"""

from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from rdetoolkit.core.flow import flow
from rdetoolkit.runner.executor import artifact_stage_order
from rdetoolkit.runner.lifecycle import Runner
from rdetoolkit.types import InputPaths, OutputContext
from tests.v2.contract.fixtures import _generate

#: ``${metadata:constant:...}`` with no metadata.json is a framework-side magic
#: failure: the flow succeeds and only artifact publication fails.
_MAGIC_TEMPLATE = "${metadata:constant:missing_key}"

_V1_ORACLE_WORKER = """
import json, os, sys
from pathlib import Path

root = Path(sys.argv[1])
mode = sys.argv[2]
sys.path.insert(0, os.getcwd())
from rdetoolkit.models.config import Config, MultiDataTileSettings, SmartTableSettings, SystemSettings
from rdetoolkit.workflows import run as v1_run

extended = {"multidatatile": "MultiDataTile", "rdeformat": "rdeformat"}.get(mode)
config = Config(
    system=SystemSettings(
        extended_mode=extended,
        save_raw=True,
        save_nonshared_raw=True,
        save_thumbnail_image=False,
        magic_variable=True,
        save_invoice_to_structured=True,
    ),
    multidata_tile=MultiDataTileSettings(ignore_errors=False),
    smarttable=SmartTableSettings(save_table_file=False),
)

os.chdir(root)
exit_code = 0
try:
    v1_run(custom_dataset_function=lambda a, b: None, config=config)
except SystemExit as error:
    exit_code = int(error.code or 0)

data_root = root / "data"
job_failed = data_root / "job.failed"
observation = {
    "exit_code": exit_code,
    "job_failed_code": (
        job_failed.read_text(encoding="utf-8").splitlines()[0] if job_failed.exists() else None
    ),
    "structured_invoices": sorted(
        path.relative_to(data_root).as_posix()
        for path in data_root.rglob("structured/invoice.json")
    ),
}
(root / ".stage_order_observation.json").write_text(json.dumps(observation), encoding="utf-8")
"""


@flow
def _noop_flow(paths: InputPaths) -> None:
    """Succeed on every tile so only artifact publication can fail."""
    assert paths.inputdata.is_dir()


def _v2_overrides(mode: str) -> dict[str, Any]:
    extended_mode = {"multidatatile": "MultiDataTile", "rdeformat": "rdeformat"}.get(mode)
    return {
        "system": {
            "extended_mode": extended_mode or "invoice",
            "save_raw": True,
            "save_nonshared_raw": True,
            "save_thumbnail_image": False,
            "magic_variable": True,
            "save_invoice_to_structured": True,
        },
        "execution": {"on_iteration_error": "fail_fast"},
        "smarttable": {"save_table_file": False},
    }


def _inject_magic_failure(mode: str, root: Path) -> None:
    """Make every tile's ``basic.dataName`` an unresolvable magic variable."""
    data_root = root / "data"
    invoice_path = data_root / "invoice" / "invoice.json"
    invoice = json.loads(invoice_path.read_text(encoding="utf-8"))
    invoice.setdefault("basic", {})["dataName"] = _MAGIC_TEMPLATE
    invoice_path.write_text(json.dumps(invoice), encoding="utf-8")
    if mode != "excelinvoice":
        return
    # The ExcelInvoice tile invoice takes dataName from the workbook row, so the
    # template has to live there instead.
    workbook_path = data_root / "inputdata" / "contract_excel_invoice.xlsx"
    workbook = load_workbook(workbook_path)
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value in {"test1", "test2"}:
                    cell.value = _MAGIC_TEMPLATE
    workbook.save(workbook_path)


def _observe_v1(mode: str, root: Path) -> dict[str, Any]:
    """Run v1 on the same prepared case in its own process."""
    _generate.materialize_sut_case(mode, root)
    _inject_magic_failure(mode, root)
    completed = subprocess.run(  # noqa: S603
        [sys.executable, "-c", _V1_ORACLE_WORKER, str(root), mode],
        cwd=_generate.REPOSITORY_ROOT,
        check=False,
        capture_output=True,
        text=True,
    )
    observation_path = root / ".stage_order_observation.json"
    if not observation_path.exists():
        message = f"v1 stage-order oracle failed for {mode}: {completed.stderr[-1500:]}"
        raise RuntimeError(message)
    result: dict[str, Any] = json.loads(observation_path.read_text(encoding="utf-8"))
    return result


def _observe_v2(mode: str, root: Path, monkeypatch: pytest.MonkeyPatch) -> dict[str, Any]:
    """Run the v2 Runner on the same prepared case."""
    _generate.materialize_sut_case(mode, root)
    _inject_magic_failure(mode, root)
    monkeypatch.chdir(root)
    runner = Runner(
        root=root,
        inputdata_path=root / "data" / "inputdata",
        unpacked_dir_path=root / "data" / "temp",
    )
    report = runner.run(_noop_flow, **_v2_overrides(mode))
    data_root = root / "data"
    job_failed = data_root / "job.failed"
    return {
        "status": report.status,
        "job_failed_code": (
            job_failed.read_text(encoding="utf-8").splitlines()[0] if job_failed.exists() else None
        ),
        "structured_invoices": sorted(
            path.relative_to(data_root).as_posix()
            for path in data_root.rglob("structured/invoice.json")
        ),
    }


@pytest.mark.parametrize(
    ("mode", "expected"),
    [
        pytest.param(
            "invoice",
            ("thumbnail", "structured", "magic", "description"),
            id="TC-IRA-5-EP-030-invoice",
        ),
        pytest.param(
            "smarttable",
            ("thumbnail", "structured", "magic", "description"),
            id="TC-IRA-5-EP-030-smarttable",
        ),
        pytest.param(
            "multidatatile",
            ("magic", "thumbnail", "structured", "description"),
            id="TC-IRA-5-EP-030-multidatatile",
        ),
        pytest.param(
            "excelinvoice",
            ("magic", "thumbnail", "structured", "description"),
            id="TC-IRA-5-EP-030-excelinvoice",
        ),
        pytest.param(
            "rdeformat",
            ("thumbnail", "description"),
            id="TC-IRA-5-EP-030-rdeformat",
        ),
    ],
)
def test_stage_order_matches_the_v1_pipeline__tc_ira_5_ep_030(
    mode: str,
    expected: tuple[str, ...],
    tmp_path: Path,
) -> None:
    """TC-IRA-5-EP-030: each mode publishes its artifacts in v1's order."""
    # Given: a plan for one mode
    from rdetoolkit.modes.install import install_default_handlers  # noqa: PLC0415
    from rdetoolkit.runner.mode_resolver import ModeKind  # noqa: PLC0415
    from tests.v2.runner.stage_order_helpers import make_plan  # noqa: PLC0415

    install_default_handlers()
    plan = make_plan(ModeKind(mode), tmp_path)

    # When: asking the executor for this mode's artifact stage sequence
    actual = artifact_stage_order(plan)

    # Then: the sequence is v1's, in order
    assert actual == expected


@pytest.mark.parametrize(
    "mode",
    [
        pytest.param("multidatatile", id="TC-IRA-5-EP-031"),
        pytest.param("excelinvoice", id="TC-IRA-5-EP-032"),
        pytest.param("invoice", id="TC-IRA-5-EP-033"),
    ],
)
def test_magic_failure_leaves_the_v1_artifacts__tc_ira_5_ep_031_032_033(
    mode: str,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """TC-IRA-5-EP-031/032/033: a failed magic step leaves v1's artifacts.

    This is review R3's counterexample. For MultiDataTile and ExcelInvoice v1
    expands the magic variable *first*, so the structured export never happens;
    for invoice mode the export precedes it and therefore survives.
    """
    # Given: v1 executed on the same input with an unresolvable magic variable
    expected = _observe_v1(mode, tmp_path / "v1")

    # When: the v2 Runner executes the same input and configuration
    actual = _observe_v2(mode, tmp_path / "v2", monkeypatch)

    # Then: the run failed with the v1 code and left exactly v1's artifacts
    assert expected["job_failed_code"] is not None
    assert actual["status"] == "failed"
    assert actual["job_failed_code"] == expected["job_failed_code"]
    assert actual["structured_invoices"] == expected["structured_invoices"]


#: ``${filename}`` resolves from the tile's first raw file, so a completed magic
#: step is visible as a concrete file name in the tile invoice.
_FILENAME_TEMPLATE = "${filename}"

_V1_STRUCTURED_FAILURE_WORKER = """
import json, os, shutil, sys
from pathlib import Path

root = Path(sys.argv[1])
mode = sys.argv[2]
sys.path.insert(0, os.getcwd())
from rdetoolkit.models.config import Config, MultiDataTileSettings, SmartTableSettings, SystemSettings
from rdetoolkit.workflows import run as v1_run

extended = {"multidatatile": "MultiDataTile", "rdeformat": "rdeformat"}.get(mode)
config = Config(
    system=SystemSettings(
        extended_mode=extended,
        save_raw=True,
        save_nonshared_raw=True,
        save_thumbnail_image=False,
        magic_variable=True,
        save_invoice_to_structured=True,
    ),
    multidata_tile=MultiDataTileSettings(ignore_errors=False),
    smarttable=SmartTableSettings(save_table_file=False),
)


def sabotage_structured(srcpaths, resource_paths):
    # The dataset callback succeeds but leaves ``structured`` as a plain file,
    # so the StructuredInvoiceSaver that follows cannot create its directory.
    shutil.rmtree(resource_paths.struct)
    resource_paths.struct.write_text("not a directory", encoding="utf-8")


os.chdir(root)
exit_code = 0
try:
    v1_run(custom_dataset_function=sabotage_structured, config=config)
except SystemExit as error:
    exit_code = int(error.code or 0)

data_root = root / "data"
observation = {
    "exit_code": exit_code,
    "job_failed": (data_root / "job.failed").exists(),
    "structured_invoices": sorted(
        path.relative_to(data_root).as_posix()
        for path in data_root.rglob("structured/invoice.json")
    ),
    "invoice_data_names": {
        path.relative_to(data_root).as_posix(): json.loads(path.read_text(encoding="utf-8"))["basic"]["dataName"]
        for path in sorted(data_root.rglob("invoice/invoice.json"))
    },
}
(root / ".structured_failure_observation.json").write_text(json.dumps(observation), encoding="utf-8")
"""


@flow
def _structured_sabotaging_flow(paths: InputPaths, out: OutputContext) -> None:
    """Succeed, but leave ``structured`` as a plain file for the stage after."""
    assert paths.inputdata.is_dir()
    shutil.rmtree(out.struct)
    out.struct.write_text("not a directory", encoding="utf-8")


def _inject_template(mode: str, root: Path, template: str) -> None:
    """Put ``template`` into every tile's ``basic.dataName``."""
    data_root = root / "data"
    invoice_path = data_root / "invoice" / "invoice.json"
    invoice = json.loads(invoice_path.read_text(encoding="utf-8"))
    invoice.setdefault("basic", {})["dataName"] = template
    invoice_path.write_text(json.dumps(invoice), encoding="utf-8")
    if mode != "excelinvoice":
        return
    workbook_path = data_root / "inputdata" / "contract_excel_invoice.xlsx"
    workbook = load_workbook(workbook_path)
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value in {"test1", "test2"}:
                    cell.value = template
    workbook.save(workbook_path)


def _observe_v1_structured_failure(mode: str, root: Path) -> dict[str, Any]:
    _generate.materialize_sut_case(mode, root)
    _inject_template(mode, root, _FILENAME_TEMPLATE)
    completed = subprocess.run(  # noqa: S603
        [sys.executable, "-c", _V1_STRUCTURED_FAILURE_WORKER, str(root), mode],
        cwd=_generate.REPOSITORY_ROOT,
        check=False,
        capture_output=True,
        text=True,
    )
    observation_path = root / ".structured_failure_observation.json"
    if not observation_path.exists():
        message = f"v1 structured-failure oracle failed for {mode}: {completed.stderr[-1500:]}"
        raise RuntimeError(message)
    result: dict[str, Any] = json.loads(observation_path.read_text(encoding="utf-8"))
    return result


def _observe_v2_structured_failure(mode: str, root: Path, monkeypatch: pytest.MonkeyPatch) -> dict[str, Any]:
    _generate.materialize_sut_case(mode, root)
    _inject_template(mode, root, _FILENAME_TEMPLATE)
    monkeypatch.chdir(root)
    runner = Runner(
        root=root,
        inputdata_path=root / "data" / "inputdata",
        unpacked_dir_path=root / "data" / "temp",
    )
    report = runner.run(_structured_sabotaging_flow, **_v2_overrides(mode))
    data_root = root / "data"
    return {
        "status": report.status,
        "error_code": (report.error or {}).get("code"),
        "job_failed": (data_root / "job.failed").exists(),
        "structured_invoices": sorted(
            path.relative_to(data_root).as_posix()
            for path in data_root.rglob("structured/invoice.json")
        ),
        "invoice_data_names": {
            path.relative_to(data_root).as_posix(): json.loads(path.read_text(encoding="utf-8"))["basic"]["dataName"]
            for path in sorted(data_root.rglob("invoice/invoice.json"))
        },
    }


@pytest.mark.parametrize(
    ("mode", "magic_completed"),
    [
        pytest.param("multidatatile", True, id="TC-IRA-5-EP-036"),
        pytest.param("excelinvoice", True, id="TC-IRA-5-EP-037"),
        pytest.param("invoice", False, id="TC-IRA-5-EP-038"),
    ],
)
def test_structured_failure_after_magic_keeps_the_expanded_invoice__tc_ira_5_ep_036_037_038(
    mode: str,
    magic_completed: bool,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """TC-IRA-5-EP-036/037/038: a later structured failure keeps an earlier magic.

    Review R3's required test, the other direction: when the *structured*
    export fails, the tile invoice must already carry the expanded magic
    variable for MultiDataTile / ExcelInvoice (magic runs first there), and
    must still carry the raw template for invoice mode (structured runs first).
    The run fails either way; v1 and v2 attribute the framework error with
    different catalog codes by contract, so only the presence of ``job.failed``
    is compared.
    """
    # Given: v1 executed with a callback that sabotages the structured directory
    expected = _observe_v1_structured_failure(mode, tmp_path / "v1")

    # When: the v2 Runner executes a flow that does the same
    actual = _observe_v2_structured_failure(mode, tmp_path / "v2", monkeypatch)

    # Then: both runs failed and left the same invoices, structured export absent
    assert expected["exit_code"] == 1
    assert expected["job_failed"] is True
    assert actual["status"] == "failed"
    assert actual["job_failed"] is True
    assert actual["error_code"] == 3006
    assert actual["structured_invoices"] == expected["structured_invoices"] == []
    assert actual["invoice_data_names"] == expected["invoice_data_names"]

    # And: whether magic completed before the failure is exactly v1's order
    data_names = set(actual["invoice_data_names"].values())
    assert data_names
    if magic_completed:
        assert _FILENAME_TEMPLATE not in data_names
    else:
        assert data_names == {_FILENAME_TEMPLATE}


def test_frozenset_seam_is_gone__tc_ira_5_ev_034() -> None:
    """TC-IRA-5-EV-034: the order-blind seam has no successor."""
    # Given / When: the mode protocol module
    import rdetoolkit.modes.protocol as module  # noqa: PLC0415

    # Then: nothing can express artifact selection without an order again
    assert not hasattr(module, "InvoiceStageProvider")
    assert not hasattr(module.ModeHandler, "invoice_stage_steps")


def test_default_order_is_the_invoice_pipeline__tc_ira_5_ev_035(tmp_path: Path) -> None:
    """TC-IRA-5-EV-035: a handler without the capability gets the v1 invoice order."""
    # Given: a plan for a mode with no registered handler
    from rdetoolkit.modes.registry import clear  # noqa: PLC0415
    from rdetoolkit.runner.mode_resolver import ModeKind  # noqa: PLC0415
    from tests.v2.runner.stage_order_helpers import make_plan  # noqa: PLC0415

    clear()
    plan = make_plan(ModeKind.invoice, tmp_path)

    # When / Then: the default sequence is the v1 invoice pipeline's
    assert artifact_stage_order(plan) == ("thumbnail", "structured", "magic", "description")
